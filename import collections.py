import collections 
import collections.abc
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

# 1. Initialize Presentation
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# 2. Color Palette
BLACK = RGBColor(0, 0, 0)
RED = RGBColor(204, 0, 0)         # Presenter name red
PURPLE = RGBColor(112, 48, 160)   # Subtitle purple
LINE_BLUE = RGBColor(31, 119, 180) # Slide divider line
GRAY = RGBColor(128, 128, 128)

# 3. Helper: Add horizontal blue separator line
def add_header_line(slide):
    line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0.5), Inches(1.2), Inches(12.333), Inches(0.02)
    )
    line.fill.solid()
    line.fill.fore_color.rgb = LINE_BLUE
    line.line.color.rgb = LINE_BLUE

# 4. Helper: Add page number in bottom-right
def add_page_number(slide, num):
    page_box = slide.shapes.add_textbox(Inches(12.0), Inches(6.8), Inches(1.0), Inches(0.4))
    p_page = page_box.text_frame.paragraphs[0]
    p_page.text = str(num)
    p_page.font.name = 'Arial'
    p_page.font.size = Pt(11)
    p_page.font.color.rgb = GRAY
    p_page.alignment = PP_ALIGN.RIGHT

# 5. Helper: Add slide title
def add_slide_title(slide, text):
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.4), Inches(12.333), Inches(0.8))
    tf = title_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.name = 'Arial'
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = BLACK
    p.alignment = PP_ALIGN.CENTER
    add_header_line(slide)

# 6. Helper: Add slide content (bullet points)
def add_slide_content(slide, bullets):
    content_box = slide.shapes.add_textbox(Inches(1.0), Inches(1.6), Inches(11.333), Inches(4.8))
    tf = content_box.text_frame
    tf.word_wrap = True
    for i, bullet in enumerate(bullets):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = "•  " + bullet
        p.font.name = 'Arial'
        p.font.size = Pt(20) # 18pt+ for readability
        p.font.color.rgb = BLACK
        p.space_after = Pt(14)


# ==================== SLIDE 1: TITLE SLIDE ====================
slide_layout = prs.slide_layouts[6] # Blank layout
slide1 = prs.slides.add_slide(slide_layout)

# Main Title
title_box = slide1.shapes.add_textbox(Inches(0.5), Inches(0.4), Inches(12.333), Inches(0.8))
p_title = title_box.text_frame.paragraphs[0]
p_title.text = "IoT Hardware Performance & Reliability Analytics"
p_title.font.name = 'Arial'
p_title.font.size = Pt(28)
p_title.font.bold = True
p_title.font.color.rgb = BLACK
p_title.alignment = PP_ALIGN.CENTER

add_header_line(slide1)

# Body Information
body_box = slide1.shapes.add_textbox(Inches(1.0), Inches(2.2), Inches(11.333), Inches(4.0))
tf_body = body_box.text_frame
tf_body.word_wrap = True

# Presenter
p_pres = tf_body.paragraphs[0]
p_pres.text = "[Your Name/s] & [Your Roll No/s]"
p_pres.font.name = 'Arial'
p_pres.font.size = Pt(20)
p_pres.font.bold = True
p_pres.font.color.rgb = RED
p_pres.alignment = PP_ALIGN.CENTER
p_pres.space_after = Pt(24)

# Industry Partner
p_ind = tf_body.add_paragraph()
p_ind.text = "Name of Industry: Hardware Operations Analytics"
p_ind.font.name = 'Arial'
p_ind.font.size = Pt(18)
p_ind.font.color.rgb = PURPLE
p_ind.alignment = PP_ALIGN.CENTER
p_ind.space_after = Pt(24)

# Mentor
p_men = tf_body.add_paragraph()
p_men.text = "Name of Institute Mentor: Self-Guided Portfolio Project"
p_men.font.name = 'Arial'
p_men.font.size = Pt(18)
p_men.font.color.rgb = PURPLE
p_men.alignment = PP_ALIGN.CENTER

# Department Footer
footer_box = slide1.shapes.add_textbox(Inches(0.5), Inches(6.5), Inches(12.333), Inches(0.6))
p_foot = footer_box.text_frame.paragraphs[0]
p_foot.text = "Department of ECE, BIT, Mesra, Ranchi-835215"
p_foot.font.name = 'Arial'
p_foot.font.size = Pt(14)
p_foot.font.bold = True
p_foot.font.color.rgb = BLACK
p_foot.alignment = PP_ALIGN.CENTER


# ==================== SLIDES 2 - 10: CONTENT SLIDES ====================
slides_data = [
    # Slide 2: Outline
    {
        "title": "Typical Presentation Flow",
        "bullets": [
            "Outline slide of your talk",
            "Introduction / Motivation / Problem or Challenge",
            "Details of work",
            "State how your results compare to other reported work",
            "Conclusion slide",
            "Backup slides if desired"
        ]
    },
    # Slide 3: Intro
    {
        "title": "Introduction & Motivation",
        "bullets": [
            "Operations manage 50 active fleet devices.",
            "Unplanned downtime hurts operational business lines.",
            "Manual hardware tracking is too slow.",
            "Objective: Predict degradation before hardware failures.",
            "Goal: Shift from reactive to proactive."
        ]
    },
    # Slide 4: SQL details
    {
        "title": "Details of Work - Relational SQL Ingestion",
        "bullets": [
            "Stored 432,000 rows of telemetry data.",
            "Optimized queries using composite primary indexes.",
            "Built rolling averages with window functions.",
            "Calculated Mean Time Between Failures (MTBF).",
            "Designed complex CTE thermal lag queries."
        ]
    },
    # Slide 5: Python details
    {
        "title": "Details of Work - Python Diagnostics",
        "bullets": [
            "Analyzed massive telemetry datasets in VS Code.",
            "Isolated operational outliers using IQR methods.",
            "Discovered global dataset correlation skew.",
            "Validated true degradation on segmented data.",
            "Ran Pearson tests confirming thermal degradation."
        ]
    },
    # Slide 6: BI details
    {
        "title": "Details of Work - BI Dashboard",
        "bullets": [
            "Created cloud-native Google Looker Studio reports.",
            "Connected summaries to separate data sources.",
            "Added dynamic regional and component dropdowns.",
            "Plotted thermal trends against error rates.",
            "Designed risk-prioritized scatter matrix."
        ]
    },
    # Slide 7: Comparison / Paradox
    {
        "title": "Analysis of Results",
        "bullets": [
            "Pooled data showed false positive correlations.",
            "Suggests voltage rises as temperature increases.",
            "Segmented data revealed true decay trends.",
            "Simpson's Paradox successfully resolved.",
            "Saved operations from false analytical paths."
        ]
    },
    # Slide 8: Findings
    {
        "title": "Key Findings",
        "bullets": [
            "Proactively flagged 3 critical failing devices.",
            "Power Modules ran hottest at 55°C.",
            "Shortest MTBF clocked at 150 hours.",
            "Set 45°C thermal threshold for warnings.",
            "Scatter plot mapped exact priority units."
        ]
    },
    # Slide 9: Conclusion
    {
        "title": "Summary & Future Work",
        "bullets": [
            "Successfully built end-to-end telemetry system.",
            "Dashboard delivers instant, actionable maintenance alerts.",
            "Next: Automate daily database query refreshes.",
            "Next: Deploy machine learning fail predictions.",
            "High-impact business-focused metrics established."
        ]
    },
    # Slide 10: Backup
    {
        "title": "SQL Optimization (Backup)",
        "bullets": [
            "Session sort buffers expanded to 16MB.",
            "Solved MySQL query timeout constraints.",
            "Avoided raw data load latency bottlenecks.",
            "Daily summary aggregations optimize dashboard load."
        ]
    }
]

# Generate Slide 2 to 10
for index, data in enumerate(slides_data):
    slide = prs.slides.add_slide(slide_layout)
    add_slide_title(slide, data["title"])
    add_slide_content(slide, data["bullets"])
    add_page_number(slide, index + 2) # Page 2 is index 0 of loop

# Save Presentation
output_file = "IoT_Project_Presentation.pptx"
prs.save(output_file)
print(f"Presentation generated successfully: {output_file}")